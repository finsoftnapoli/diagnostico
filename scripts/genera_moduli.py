"""Genera src/lib/modules.ts e il file Excel di revisione dai contenuti in scripts/moduli_content.py.
Uso: python3 scripts/genera_moduli.py [cartella per l'Excel]  (dalla cartella del progetto)"""
import json, pathlib, sys
sys.path.insert(0, str(pathlib.Path(__file__).parent))
import moduli_content as M
import content as C

ROOT = pathlib.Path(__file__).parent.parent
NAMES = {a[0]: a[1] for a in C.AMBITI}

subareas = {}
modules = {}
for area, subs in M.MODULI.items():
    assert area in NAMES, area
    assert len(subs) == 8, (area, "servono 8 sotto-ambiti", len(subs))
    subareas[area] = []
    out = []
    for si, (sname, sdesc, qs) in enumerate(subs, start=1):
        scode = f"{area}-{si}"
        assert 2 <= len(qs) <= 3, (scode, "servono 2-3 domande", len(qs))
        short = M.BREVI.get(sname)
        assert short and len(short) <= 24, (scode, sname, "manca o troppo lunga l'etichetta breve")
        subareas[area].append(dict(code=scode, name=sname, short=short, description=sdesc))
        for qi, qq in enumerate(qs, start=1):
            qid = f"{scode}.{qi}"
            levels = qq["levels"]
            assert len(levels) == 4 and all(l.strip() for l in levels), qid
            chk = None
            if qq["check"]:
                ctext, opts, min_level, flag = qq["check"]
                assert all(x in opts for x in flag), (qid, flag, opts)
                chk = dict(text=ctext, options=opts, flag=dict(minLevel=min_level, answers=flag))
            out.append(dict(id=qid, area=area, sub=scode, theme=qq["theme"], text=qq["text"], levels=levels, check=chk, na=qq["na"]))
    assert len({q["id"] for q in out}) == len(out), area
    modules[area] = out

ts = f"""// Generato da scripts/genera_moduli.py (moduli di approfondimento v{M.VERSION}). Non modificare a mano:
// cambiare scripts/moduli_content.py e rigenerare.
import type {{ AreaCode }} from "./types";

export const MODULE_VERSION = "{M.VERSION}";

/** Gli 8 sotto-ambiti di ciascun modulo: i vertici dell'Ottogramma dell'ambito. */
export interface ModuleSubarea {{
  code: string;
  name: string;
  /** Etichetta breve per i vertici dell'Ottogramma. */
  short: string;
  description: string;
}}

export interface ModuleQuestion {{
  id: string;
  area: AreaCode;
  /** Codice del sotto-ambito a cui appartiene la domanda. */
  sub: string;
  theme: string;
  text: string;
  levels: string[];
  check: {{ text: string; options: string[]; flag: {{ minLevel: number; answers: string[] }} }} | null;
  /** La domanda può non riguardare l'azienda: si può rispondere "non applicabile". */
  na: boolean;
}}

export const MODULE_SUBAREAS: Record<AreaCode, ModuleSubarea[]> = {json.dumps(subareas, ensure_ascii=False, indent=2)};

export const MODULE_QUESTIONS: Record<AreaCode, ModuleQuestion[]> = {json.dumps(modules, ensure_ascii=False, indent=2)};

export const ROADMAP_HORIZONS = {json.dumps(M.ORIZZONTI, ensure_ascii=False)};
"""
(ROOT / "src/lib/modules.ts").write_text(ts, encoding="utf-8")
print("src/lib/modules.ts:", sum(len(v) for v in modules.values()), "domande,", sum(len(v) for v in subareas.values()), "sotto-ambiti")

# ---------- Excel di revisione ----------
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

NAVY = "0A2A4A"
HDR = PatternFill("solid", fgColor=NAVY)
SUBFILL = PatternFill("solid", fgColor="E8ECF1")
NOTE = PatternFill("solid", fgColor="FFF2CC")
thin = Side(style="thin", color="BFC5CC")
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
WRAP = Alignment(wrap_text=True, vertical="top")
f = lambda **k: Font(name="Arial", size=k.pop("size", 10), **k)

wb = Workbook()
ws = wb.active
ws.title = "Leggimi"
ws["A1"] = f"Moduli di approfondimento per ambito — bozza v{M.VERSION} da rivedere"
ws["A1"].font = f(bold=True, size=14, color=NAVY)
ws["A2"] = "Testi originali. Annota le modifiche nella colonna gialla «Note di revisione» di ogni foglio: le riporto nel sistema."
ws["A2"].font = f(italic=True, color="595959")
rows = [
    ("Come funziona", "Ogni modulo si attiva con l'acquisto dell'approfondimento dell'ambito (€100 + IVA) o con il percorso completo."),
    ("Struttura", "Come il diagnostico, ogni modulo è un Ottogramma dell'ambito: 8 sotto-ambiti, ciascuno con 2 o 3 domande. Ogni domanda ha 4 livelli di maturità dal più basso al più alto."),
    ("Verifica", "Alcune domande hanno una domanda di verifica: se il livello scelto è alto ma la verifica dice il contrario, il report segnala una possibile sovrastima."),
    ("Non applicabile", "Le domande segnate «Sì» nella colonna N/A ammettono la risposta «non riguarda la nostra azienda» (per esempio magazzino nelle aziende di servizi)."),
    ("Punteggio", "Livello 1 = 0, 2 = 33, 3 = 67, 4 = 100. Sotto-ambito = media delle sue domande applicabili; modulo = media degli 8 sotto-ambiti, come nel diagnostico."),
    ("Report", "Ottogramma dell'ambito subito dopo l'invio; poi l'AI scrive diagnosi per sotto-ambito, piano a 3, 6 e 12 mesi e indicatori, rivisti dallo staff prima dell'invio."),
]
ws.column_dimensions["A"].width = 18
ws.column_dimensions["B"].width = 110
for i, (k, v) in enumerate(rows, start=4):
    ws.cell(row=i, column=1, value=k).font = f(bold=True)
    c = ws.cell(row=i, column=2, value=v)
    c.font = f()
    c.alignment = WRAP
r = 4 + len(rows) + 1
for col, txt in ((1, "Ambito"), (2, "Sotto-ambiti (domande)")):
    c = ws.cell(row=r, column=col, value=txt)
    c.font = f(bold=True, color="FFFFFF")
    c.fill = HDR
for area in modules:
    r += 1
    ws.cell(row=r, column=1, value=NAMES[area]).font = f()
    counts = {s["code"]: sum(1 for q in modules[area] if q["sub"] == s["code"]) for s in subareas[area]}
    c = ws.cell(row=r, column=2, value=f"{len(modules[area])} domande · " + "; ".join(f"{s['name']} ({counts[s['code']]})" for s in subareas[area]))
    c.font = f()
    c.alignment = WRAP

cols = ["ID", "Tema", "Domanda", "Livello 1", "Livello 2", "Livello 3", "Livello 4", "Domanda di verifica", "Risposte", "Segnale di sovrastima", "N/A", "Note di revisione"]
widths = [9, 20, 36, 30, 30, 30, 30, 30, 16, 22, 6, 30]
for area, qs in modules.items():
    s = wb.create_sheet(area + " " + NAMES[area].split(" ")[0])
    s.freeze_panes = "C2"
    for i, (h, w) in enumerate(zip(cols, widths), start=1):
        c = s.cell(row=1, column=i, value=h)
        c.font = f(bold=True, color="FFFFFF")
        c.fill = HDR
        c.alignment = Alignment(wrap_text=True, vertical="center")
        c.border = BOX
        s.column_dimensions[c.column_letter].width = w
    j = 1
    for sub in subareas[area]:
        j += 1
        c = s.cell(row=j, column=1, value=sub["code"])
        c.font = f(bold=True, color=NAVY)
        c.fill = SUBFILL
        c2 = s.cell(row=j, column=2, value=f"{sub['name']} — {sub['description']}")
        c2.font = f(bold=True, color=NAVY)
        s.merge_cells(start_row=j, start_column=2, end_row=j, end_column=len(cols) - 1)
        c2.fill = SUBFILL
        s.cell(row=j, column=len(cols)).fill = NOTE
        for q in (x for x in qs if x["sub"] == sub["code"]):
            j += 1
            chk = q["check"]
            vals = [q["id"], q["theme"], q["text"], *q["levels"],
                    chk["text"] if chk else "", " / ".join(chk["options"]) if chk else "",
                    (f"Livello {chk['flag']['minLevel']} o più con risposta " + " o ".join(chk["flag"]["answers"])) if chk else "",
                    "Sì" if q["na"] else "", ""]
            for i, v in enumerate(vals, start=1):
                c = s.cell(row=j, column=i, value=v)
                c.font = f()
                c.alignment = WRAP
                c.border = BOX
                if i == len(vals):
                    c.fill = NOTE

out = pathlib.Path(sys.argv[1] if len(sys.argv) > 1 else ".") / f"Moduli_approfondimento_v{M.VERSION}.xlsx"
wb.save(out)
print("Excel:", out)
